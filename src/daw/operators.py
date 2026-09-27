"""Vetted numerical operators. Native measurements stay stratified by representation."""
import csv
import math
from itertools import islice

from daw.inspectors import inspect_hdf
from daw.util import DawError, read_json


def table_rows(ws, asset, curation):
    settings = curation.settings
    first, last = curation.selector.get("start_row", 2), curation.selector.get("end_row")
    if not isinstance(first, int) or first < 1 or last is not None and (not isinstance(last, int) or last < first):
        raise DawError("invalid_table_region")
    column = settings.get("feature_column")
    columns = settings.get("columns", {})
    if not isinstance(column, int) or isinstance(column, bool) or column < 1:
        raise DawError("invalid_feature_column", "columns use one-based source coordinates")
    if any(not isinstance(c, int) or isinstance(c, bool) or c < 1 for c in columns.values()):
        raise DawError("invalid_value_column")
    path = ws.blob_path(asset["blob"])
    is_workbook = asset["body"]["name"].lower().endswith((".xlsx", ".xlsm"))
    workbook, cached, handle = None, None, None
    count = 0
    try:
        if is_workbook:
            import openpyxl
            from daw.inspectors import archive_inventory
            archive_inventory(path, ws.budgets)
            workbook = openpyxl.load_workbook(path.open("rb"), read_only=True, data_only=False, keep_links=False)
            cached = openpyxl.load_workbook(path.open("rb"), read_only=True, data_only=True, keep_links=False)
            sheet = curation.selector.get("sheet")
            if sheet not in workbook.sheetnames:
                raise DawError("unknown_sheet")
            if last is None or last > workbook[sheet].max_row:
                raise DawError("invalid_table_region", "workbooks require explicit bounded end_row")
            rows = zip(workbook[sheet].iter_rows(min_row=first, max_row=last),
                       cached[sheet].iter_rows(min_row=first, max_row=last), strict=True)
        else:
            handle = path.open(newline="", encoding="utf-8-sig")
            reader = csv.reader(handle, delimiter=settings.get("delimiter", "," if asset["body"]["name"].endswith(".csv") else "\t"))
            rows = ((row, None) for row in islice(reader, first - 1, last))
        for number, (row, cache) in enumerate(rows, first):
            count += 1
            if count > ws.budgets.max_rows or count * len(row) > ws.budgets.max_cells:
                raise DawError("over_budget", "table region scan")
            if max([column] + list(columns.values())) > len(row):
                raise DawError("column_outside_region", f"row {number}")
            feature_cell = row[column - 1]
            feature = feature_cell.value if is_workbook else feature_cell
            if feature is None or feature == "":
                continue
            if not isinstance(feature, str):
                raise DawError("nonliteral_feature_identifier", f"row {number}: identifiers must be source text")
            values, unresolved = {}, []
            for key, col in columns.items():
                cell = row[col - 1]
                if is_workbook and cell.data_type == "f":
                    value = cache[col - 1].value
                    if value is None:
                        unresolved.append(key)
                else:
                    value = cell.value if is_workbook else cell
                values[key] = value
            locator = f"sheet:{curation.selector['sheet']}!row:{number}" if is_workbook else f"row:{number}"
            yield {"feature": feature, "source_row": number, "locator": locator,
                   "feature_column": column, "columns": columns, "values": values, "unresolved": unresolved}
    finally:
        for obj in (workbook, cached, handle):
            if obj:
                obj.close()


def table_lookup(ws, asset, curation, request):
    matches = []
    for row in table_rows(ws, asset, curation):
        if row["feature"] == request.feature:
            for field, value in row["values"].items():
                numeric = None
                if value not in (None, ""):
                    try:
                        numeric = float(value)
                    except (ValueError, TypeError):
                        pass
                if numeric is not None and not math.isfinite(numeric):
                    numeric = None
                matches.append({"feature": request.feature, "field": field, "value": numeric,
                                "text": None if value is None else str(value),
                                "locator": row["locator"] + f":column:{row['columns'][field]}",
                                "source_row": row["source_row"], "criterion": "literal source row found; no significance test",
                                "units": curation.settings.get("column_units", {}).get(field, curation.settings.get("units", "unspecified")),
                                "value_state": "unresolved_formula_cache" if field in row["unresolved"] else "reported"})
            if not row["values"]:
                matches.append({"feature": request.feature, "locator": row["locator"], "text": request.feature,
                                "criterion": "literal source row found"})
    if matches:
        return matches, "evaluated_detected", "literal_row_present_not_a_significance_claim"
    selection = curation.settings.get("selection", "unknown")
    if selection in {"significant_only", "selected", "unknown"}:
        return [], "unresolved", "absent_from_selected_or_unknown_universe; no_null_effect_inference"
    return [], "not_measurable", "feature_not_in_declared_table_universe"


def checked_region(ws, curation, region):
    if region.reference != curation.reference:
        raise DawError("incompatible_reference", "native assembly queries only; no implicit liftover")
    reference = read_json(ws.blob_path(curation.reference))
    chrom = reference["aliases"].get(region.chrom, region.chrom)
    if chrom not in reference["chromosomes"] or region.end > reference["chromosomes"][chrom]:
        raise DawError("region_outside_reference")
    return chrom, region.start, region.end, reference


def interval_overlap(ws, asset, curation, request):
    import bioframe
    import pandas as pd

    chrom, start, end, reference = checked_region(ws, curation, request.region)
    convention = curation.settings["native_coordinates"]
    if convention not in {"0-based-half-open", "1-based-closed"}:
        raise DawError("unknown_coordinate_convention")
    schema = curation.settings["interval_schema"]
    if schema not in {"bed3", "bed4", "bed6", "narrowPeak", "broadPeak"}:
        raise DawError("unsupported_interval_schema")
    required = {"bed3": 3, "bed4": 4, "bed6": 6, "narrowPeak": 10, "broadPeak": 9}[schema]
    rows = []
    with ws.blob_path(asset["blob"]).open() as f:
        for line_number, line in enumerate(f, 1):
            if line.startswith(("#", "track", "browser")) or not line.strip():
                continue
            if len(rows) >= ws.budgets.max_rows:
                raise DawError("over_budget", "interval row budget")
            fields = line.rstrip("\n").split("\t")
            if len(fields) < required:
                raise DawError("invalid_interval_schema", f"line {line_number}")
            source_start, source_end = int(fields[1]), int(fields[2])
            native_chrom = reference["aliases"].get(fields[0], fields[0])
            left = source_start - (convention == "1-based-closed")
            if left < 0 or source_end <= left or native_chrom not in reference["chromosomes"] or source_end > reference["chromosomes"][native_chrom]:
                raise DawError("interval_outside_reference", f"line {line_number}")
            rows.append({"chrom": native_chrom, "start": left, "end": source_end, "source_row": line_number,
                         "source_start": source_start, "source_end": source_end, "text": "\t".join(fields[3:])})
    if not rows:
        return [], "evaluated_not_detected", "no_called_peak_in_file_under_source_processing; not_absence_of_occupancy"
    frame = pd.DataFrame(rows)
    selected = bioframe.overlap(frame, pd.DataFrame([{"chrom": chrom, "start": start, "end": end}]), how="inner", suffixes=("", "_query"))
    result = []
    for row in selected.to_dict("records"):
        result.append({k: v for k, v in row.items() if not k.endswith("_query")})
        result[-1]["locator"] = f"line:{row['source_row']}"
        result[-1]["criterion"] = "positive-length half-open overlap of source-called intervals"
    return result, "evaluated_detected" if result else "evaluated_not_detected", "called_peak_overlap; not_molecular_occupancy_test"


def signal_summary(ws, asset, curation, request):
    import pyBigWig

    chrom, start, end, reference = checked_region(ws, curation, request.region)
    if curation.settings["signal_missingness"] != "covered_bases_only":
        raise DawError("unsupported_missingness_policy", "zero fill needs a separately reviewed operator")
    path = ws.blob_path(asset["blob"])
    with pyBigWig.open(str(path)) as bw:
        if not bw.isBigWig() or bw.chroms().get(chrom) != reference["chromosomes"][chrom]:
            raise DawError("signal_reference_mismatch")
        intervals = bw.intervals(chrom, start, end) or []
        if len(intervals) > ws.budgets.max_rows:
            raise DawError("over_budget", "signal intervals")
        covered, weighted = 0, 0.0
        for left, right, value in intervals:
            if not math.isfinite(value):
                continue
            length = min(end, right) - max(start, left)
            if length > 0:
                covered += length
                weighted += length * value
        exact = bw.stats(chrom, start, end, type="mean", exact=True)[0]
    value = weighted / covered if covered else None
    if value is not None and exact is not None and not math.isclose(value, exact, rel_tol=1e-6, abs_tol=1e-8):
        raise DawError("signal_validation_failed", "exact summary differs from weighted intervals")
    row = {"chrom": chrom, "start": start, "end": end, "value": value, "field": "covered_mean",
           "requested_bases": end - start, "covered_bases": covered, "coverage_fraction": covered / (end - start),
           "statistic": "exact mean over covered bases", "units": curation.settings["units"],
           "missing_policy": "uncovered bases excluded, never implicitly zero", "binning": "native intervals",
           "locator": f"{chrom}:{start}-{end}", "criterion": "descriptive extraction; no significance threshold"}
    return [row], "evaluated_detected" if covered else "not_measurable", "signal_extracted" if covered else "no_covered_bases"


def matrix_values(ws, asset, curation, request):
    import numpy as np
    from scipy import sparse

    path = ws.blob_path(asset["blob"])
    settings = curation.settings
    data, close, symbols = None, None, None
    if asset["body"]["name"].endswith(".mtx"):
        from scipy.io import mmread
        required = {"feature_ids", "cell_ids", "bundle_identity", "orientation"}
        if not required <= settings.keys():
            raise DawError("unpaired_matrix_market_bundle")
        identities = settings["bundle_identity"]
        if not isinstance(identities, dict) or set(identities) != {"matrix", "features", "cells"} or len(set(identities.values())) != 1:
            raise DawError("matrix_bundle_release_mismatch")
        feature_ids = read_json(ws.blob_path(settings["feature_ids"]))
        cell_ids = read_json(ws.blob_path(settings["cell_ids"]))
        from daw.inspectors import inspect_file
        info = inspect_file(path, "matrix.mtx", ws.budgets)
        dims = info.get("dimensions", [])
        if len(dims) != 3 or dims[2] * 32 > ws.budgets.worker_memory_bytes:
            raise DawError("over_budget", "Matrix Market sparse allocation")
        data = mmread(path).tocsr()
        if settings["orientation"] == "features_by_cells":
            data = data.T.tocsr()
        elif settings["orientation"] != "cells_by_features":
            raise DawError("unknown_matrix_orientation")
        if data.shape != (len(cell_ids), len(feature_ids)):
            raise DawError("matrix_axis_mismatch")
    else:
        import anndata
        structure = inspect_hdf(path, ws.budgets)
        # AnnData backed mode still eagerly loads layers and axis metadata. Bound
        # their expanded footprint before invoking the high-level reader.
        eager_bytes = 0
        for obj in structure["objects"]:
            if "shape" in obj and not obj["path"].startswith(("/X/", "/raw/X/")) and obj["path"] not in {"/X", "/raw/X"}:
                eager_bytes += math.prod(obj["shape"]) * max(obj.get("itemsize", 64), 64 if obj.get("variable_length") else 1)
        if eager_bytes * 3 > ws.budgets.worker_memory_bytes:
            raise DawError("over_budget", "AnnData eager layers/metadata plus temporary copies")
        ad = anndata.read_h5ad(path, backed="r")
        close = ad.file.close
        cell_ids = list(map(str, ad.obs_names))
        layer = settings["layer"]
        if layer == "raw.X":
            if ad.raw is None:
                close()
                raise DawError("missing_matrix_layer")
            data, feature_ids = ad.raw.X, list(map(str, ad.raw.var_names))
            feature_frame = ad.raw.var
        elif layer == "X":
            data, feature_ids = ad.X, list(map(str, ad.var_names))
            feature_frame = ad.var
        elif layer.startswith("layers/") and layer[7:] in ad.layers:
            data, feature_ids = ad.layers[layer[7:]], list(map(str, ad.var_names))
            feature_frame = ad.var
        else:
            close()
            raise DawError("missing_matrix_layer")
        id_column = settings.get("feature_id_column", "index")
        symbol_column = settings.get("gene_symbol_column")
        if symbol_column:
            if symbol_column != "index" and symbol_column not in feature_frame:
                close()
                raise DawError("missing_gene_symbol_column")
            symbols = list(map(str, feature_frame.index if symbol_column == "index" else feature_frame[symbol_column]))
        if id_column != "index":
            if id_column not in feature_frame:
                close()
                raise DawError("missing_feature_identifier_column")
            feature_ids = feature_frame[id_column].tolist()
            if any(not isinstance(x, str) for x in feature_ids):
                close()
                raise DawError("nonliteral_feature_identifier")
    try:
        if len(cell_ids) != len(set(cell_ids)) or len(feature_ids) != len(set(feature_ids)):
            raise DawError("duplicate_axis_identifiers", "gene symbols must be separate from feature IDs")
        if set(curation.sample_map) != set(cell_ids):
            raise DawError("cell_mapping_mismatch", "explicit map must cover exactly the selected matrix cells")
        if request.feature not in feature_ids:
            return [], "not_measurable", "feature_absent_from_selected_layer"
        if len(cell_ids) * 32 > ws.budgets.worker_memory_bytes or len(cell_ids) > ws.budgets.max_cells:
            raise DawError("over_budget", "feature slice allocation")
        feature_index = feature_ids.index(request.feature)
        values = data[:, feature_index:feature_index + 1]
        if sparse.issparse(values):
            values = values.toarray()
        values = np.asarray(values).ravel()
        presence = settings.get("feature_presence", {})
        rows, grouped = [], {}
        for i, cell in enumerate(cell_ids):
            metadata = curation.sample_map[cell]
            dataset = metadata.get("dataset", metadata.get("study", ""))
            measured = request.feature in presence.get(dataset, []) if presence else False
            val = float(values[i])
            if not measured:
                rows.append({"feature": request.feature, "cell": cell, "value": None, "value_state": "feature_presence_unresolved",
                             "locator": f"{settings['layer']}[{i},{feature_index}]"})
                continue
            if not math.isfinite(val):
                raise DawError("nonfinite_matrix_value")
            if request.operator == "expression.pseudobulk_counts":
                if val < 0 or not val.is_integer():
                    raise DawError("counts_validation_failed", "no rounding or conversion of normalized data")
                key = (metadata["study"], metadata["donor"], metadata["sample"])
                acc = grouped.setdefault(key, {"sum": 0, "cells": 0})
                acc["sum"] += int(val)
                acc["cells"] += 1
            else:
                rows.append({"feature": request.feature, "cell": cell, "sample": metadata.get("sample"),
                             "gene_symbol": symbols[feature_index] if symbols is not None else None,
                             "donor": metadata.get("donor"), "study": metadata.get("study"), "value": val,
                             "value_state": "measured", "locator": f"{settings['layer']}[{i},{feature_index}]",
                             "criterion": "measured feature value; zero retained without detection inference"})
        if request.operator == "expression.pseudobulk_counts":
            if rows:
                return rows, "unresolved", "feature_presence_required_for_every_grouped_cell"
            rows = [{"feature": request.feature, "study": key[0], "donor": key[1], "sample": key[2],
                     "value": value["sum"], "cell_count": value["cells"], "field": "pseudobulk_count",
                     "criterion": "sum of source counts by documented biological sample; no differential test"}
                    for key, value in sorted(grouped.items())]
        unresolved = any(r.get("value_state") == "feature_presence_unresolved" for r in rows)
        return rows, "unresolved" if unresolved else "evaluated_detected", "presence_unresolved" if unresolved else "descriptive_expression_values"
    finally:
        if close:
            close()


def contact_extract(ws, asset, curation, request):
    import cooler
    import numpy as np

    path = ws.blob_path(asset["blob"])
    inspect_hdf(path, ws.budgets)
    chrom, start, end, reference = checked_region(ws, curation, request.region)
    other = request.region2 or request.region
    chrom2, start2, end2, _ = checked_region(ws, curation, other)
    group = curation.selector.get("group", "/")
    cool = cooler.Cooler(str(path) + "::" + group)
    if int(cool.chromsizes.get(chrom, 0)) != reference["chromosomes"][chrom] or int(cool.chromsizes.get(chrom2, 0)) != reference["chromosomes"][chrom2]:
        raise DawError("contact_reference_mismatch")
    resolution = curation.settings["resolution"]
    if cool.binsize != resolution:
        raise DawError("contact_resolution_mismatch")
    requested = request.parameters.get("required_resolution")
    if requested is not None and resolution > requested:
        raise DawError("insufficient_contact_resolution", "requested interpretation is finer than measured bins")
    balance = curation.settings["balance"]
    if balance is not False and balance not in {True, "weight"}:
        raise DawError("unsupported_contact_normalization")
    if balance and "weight" not in cool.bins().columns:
        raise DawError("missing_balancing_weights")
    region1, region2 = (chrom, start, end), (chrom2, start2, end2)
    b1, b2 = cool.bins().fetch(region1), cool.bins().fetch(region2)
    if len(b1) * len(b2) > ws.budgets.max_cells:
        raise DawError("over_budget", "contact selection matrix")
    matrix = cool.matrix(balance=balance, sparse=True).fetch(region1, region2).tocoo()
    rows, seen = [], set()
    for i, j, value in zip(matrix.row, matrix.col, matrix.data, strict=True):
        left, right = b1.iloc[i], b2.iloc[j]
        pair = tuple(sorted((int(b1.index[i]), int(b2.index[j])))) if cool.storage_mode == "symmetric-upper" else (int(b1.index[i]), int(b2.index[j]))
        if pair in seen:
            continue
        seen.add(pair)
        invalid = not np.isfinite(value)
        rows.append({"chrom": str(left.chrom), "start": int(left.start), "end": int(left.end),
                     "chrom2": str(right.chrom), "start2": int(right.start), "end2": int(right.end),
                     "value": None if invalid else float(value), "invalid_bin": bool(invalid),
                     "resolution": resolution, "normalization": "balanced" if balance else "raw_counts",
                     "locator": f"{group}:bins:{pair[0]},{pair[1]}", "criterion": "measured contact value; not a called loop"})
    # Preserve invalid masks even for bins with no explicit stored pixels.
    mask = {"left": [int(b1.index[i]) for i in range(len(b1)) if balance and not np.isfinite(b1.iloc[i]["weight"])],
            "right": [int(b2.index[i]) for i in range(len(b2)) if balance and not np.isfinite(b2.iloc[i]["weight"])]}
    rows.append({"field": "selection_context", "text": __import__("json").dumps({"invalid_bins": mask,
                 "selected_bins": [len(b1), len(b2)], "implicit_valid_pixels": "zero", "symmetric_duplicates": "removed"}),
                 "resolution": resolution, "normalization": "balanced" if balance else "raw_counts"})
    return rows, "evaluated_detected", "contact_region_extracted; no_loop_or_promoter_specific_claim"


OPERATORS = {"table.feature_lookup": table_lookup, "contrast.published_lookup": table_lookup,
             "interval.overlap": interval_overlap, "signal.interval_summary": signal_summary,
             "expression.feature_values": matrix_values, "expression.pseudobulk_counts": matrix_values,
             "contact.region_extract": contact_extract}
