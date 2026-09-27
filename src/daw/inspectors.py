"""Structural inventories only; parsing never confers scientific eligibility."""
import csv
import datetime
import gzip
import json
import os
import subprocess
import sys
import tarfile
import uuid
import zipfile
from pathlib import Path, PurePosixPath

from defusedxml import ElementTree as ET

from daw.models import Asset, Budgets
from daw.util import DawError, file_hash, now, read_json, write_json


def safe_member(name):
    path = PurePosixPath(name.replace("\\", "/"))
    return bool(name) and not path.is_absolute() and ".." not in path.parts and ":" not in name and "\x00" not in name


def archive_inventory(path, budgets):
    members, names, total = [], set(), 0
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            if len(z.infolist()) > budgets.archive_members:
                raise DawError("over_budget", "archive member count")
            for member in z.infolist():
                name = member.filename
                mode = (member.external_attr >> 16) & 0o170000
                if not safe_member(name) or name in names or mode == 0o120000:
                    raise DawError("unsafe_to_inspect", "unsafe path, duplicate member, or symlink")
                names.add(name)
                total += member.file_size
                if total > budgets.expanded_bytes:
                    raise DawError("over_budget", "archive expansion limit")
                members.append({"name": name, "size": member.file_size, "compressed_size": member.compress_size,
                                "compression": member.compress_type, "directory": member.is_dir()})
    else:
        with tarfile.open(path, mode="r|*") as tar:
            for member in tar:
                if len(members) >= budgets.archive_members:
                    raise DawError("over_budget", "archive member count")
                if not safe_member(member.name) or member.name in names or not (member.isfile() or member.isdir()):
                    raise DawError("unsafe_to_inspect", "unsafe path, duplicate member, or nonregular TAR entry")
                names.add(member.name)
                total += member.size
                if total > budgets.expanded_bytes:
                    raise DawError("over_budget", "archive expansion limit")
                members.append({"name": member.name, "size": member.size, "directory": member.isdir()})
    return {"kind": "archive", "members": members, "expanded_bytes": total, "complete": True}


def cell_value(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return {"value": value.isoformat(), "type": "date", "warning": "source date; no gene-symbol repair"}
    return value


def inspect_workbook(path, budgets):
    import openpyxl

    archive_inventory(path, budgets)
    path = Path(path)
    formula = openpyxl.load_workbook(path.open("rb"), read_only=True, data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(path.open("rb"), read_only=True, data_only=True, keep_links=False)
    sheets, used = [], 0
    try:
        for sheet in formula:
            preview, formulas, cell_types, scanned = [], [], {}, 0
            complete = True
            for row, cache_row in zip(sheet.iter_rows(), cached[sheet.title].iter_rows(), strict=True):
                if used + len(row) > budgets.max_cells or scanned >= budgets.max_rows:
                    complete = False
                    break
                used += len(row)
                scanned += 1
                values = []
                for cell, cache in zip(row, cache_row, strict=True):
                    typ = getattr(cell, "data_type", "empty")
                    cell_types[typ] = cell_types.get(typ, 0) + 1
                    if cell.value is not None:
                        values.append({"cell": cell.coordinate, "value": cell_value(cell.value), "type": typ})
                    if typ == "f" and len(formulas) < 100:
                        formulas.append({"cell": cell.coordinate, "formula": cell.value,
                                         "cached": cell_value(cache.value), "cache_resolved": cache.value is not None})
                if values and len(preview) < budgets.preview_rows:
                    preview.append(values)
            sheets.append({"name": sheet.title, "state": sheet.sheet_state, "rows": sheet.max_row,
                           "columns": sheet.max_column, "preview": preview, "formulas": formulas,
                           "cell_types": cell_types, "scanned_rows": scanned, "complete": complete})
        # Merged-cell metadata is not exposed by openpyxl's streaming worksheet.
        with zipfile.ZipFile(path) as z:
            for sheet, xml in zip(sheets, [f"xl/worksheets/sheet{i + 1}.xml" for i in range(len(sheets))], strict=True):
                sheet["merged_ranges"] = []
                if xml in z.namelist():
                    with z.open(xml) as stream:
                        for _, elem in ET.iterparse(stream, events=("end",)):
                            if elem.tag.endswith("}mergeCell"):
                                sheet["merged_ranges"].append(elem.attrib.get("ref"))
                            elem.clear()
        return {"kind": "workbook", "sheets": sheets, "complete": all(s["complete"] for s in sheets),
                "formula_policy": "not evaluated; absent caches stay unresolved"}
    finally:
        formula.close()
        cached.close()


def inspect_table(path, budgets, delimiter="\t"):
    preview, rows, columns, complete = [], 0, set(), True
    with Path(path).open(newline="", encoding="utf-8-sig") as f:
        for row in csv.reader(f, delimiter=delimiter):
            rows += 1
            columns.add(len(row))
            if rows > budgets.max_rows or rows * max(columns) > budgets.max_cells:
                complete = False
                break
            if len(preview) < budgets.preview_rows:
                preview.append({"row": rows, "values": row})
    return {"kind": "delimited", "delimiter": delimiter, "rows_scanned": rows,
            "column_counts": sorted(columns), "preview": preview, "complete": complete,
            "missing_tokens": [], "identifier_policy": "literal strings"}


def inspect_hdf(path, budgets):
    import h5py

    objects, visited = [], set()
    with h5py.File(path, "r") as f:
        def walk(group, prefix=""):
            address = h5py.h5o.get_info(group.id).addr
            if address in visited:
                return
            visited.add(address)
            for name in group:
                if len(objects) >= budgets.archive_members:
                    raise DawError("over_budget", "HDF5 object count")
                link = group.get(name, getlink=True)
                if isinstance(link, (h5py.ExternalLink, h5py.SoftLink)):
                    raise DawError("unsafe_to_inspect", "HDF5 external/soft links are not dereferenced")
                obj = group[name]
                p = f"{prefix}/{name}"
                if isinstance(obj, h5py.Dataset):
                    if obj.is_virtual or obj.external:
                        raise DawError("unsafe_to_inspect", "external or virtual HDF5 storage")
                    objects.append({"path": p, "shape": list(obj.shape), "dtype": str(obj.dtype),
                                    "itemsize": obj.dtype.itemsize, "variable_length": h5py.check_dtype(vlen=obj.dtype) is not None,
                                    "compression": obj.compression})
                else:
                    objects.append({"path": p, "kind": "group",
                                    "encoding_type": str(obj.attrs.get("encoding-type", ""))})
                    walk(obj, p)
        walk(f)
    return {"kind": "hdf5", "objects": objects, "complete": True,
            "warning": "X/raw/layer names do not establish count semantics"}


def inspect_file(path, name, budgets=None):
    budgets = budgets or Budgets()
    suffix = name.lower().rsplit(".", 1)[-1]
    try:
        if suffix in {"pkl", "pickle", "py", "r", "sh", "rds", "rda", "rdata", "pdf", "xls", "xlsb", "ods"}:
            return {"status": "unsupported_format", "name": name,
                    "reason": "untrusted code/serialization is never executed; specialized reader not activated"}
        if suffix in {"xlsx", "xlsm"}:
            output = inspect_workbook(path, budgets)
        elif suffix in {"tsv", "csv", "txt", "bed", "narrowpeak", "broadpeak", "gtf", "gff", "gff3"}:
            output = inspect_table(path, budgets, "," if suffix == "csv" else "\t")
        elif suffix in {"h5ad", "h5", "cool", "mcool"}:
            output = inspect_hdf(path, budgets)
            if suffix in {"cool", "mcool"}:
                import cooler
                output["cooler_groups"] = cooler.fileops.list_coolers(str(path))
        elif suffix in {"bw", "bigwig", "bb", "bigbed"}:
            import pyBigWig
            with pyBigWig.open(str(path)) as big:
                output = {"kind": "bigwig" if big.isBigWig() else "bigbed", "chromosomes": big.chroms(),
                          "header": big.header(), "complete": True, "scope": "header only"}
        elif suffix in {"zip", "tar", "tgz"} or name.lower().endswith(".tar.gz"):
            output = archive_inventory(path, budgets)
        elif suffix == "gz":
            output = {"kind": "gzip", "members": [{"name": name[:-3], "size": None}], "complete": False,
                      "reason": "stream decompression requires selected-member extraction"}
        elif suffix == "mtx":
            with Path(path).open() as f:
                banner = f.readline().strip()
                for line in f:
                    if not line.startswith("%"):
                        dimensions = [int(x) for x in line.split()]
                        break
                else:
                    raise ValueError("missing dimensions")
            output = {"kind": "matrix_market", "banner": banner, "dimensions": dimensions, "complete": False,
                      "reason": "paired feature/barcode files and release identity require curation"}
        elif suffix in {"json", "xml"}:
            # Metadata is inventoried but not transformed into scientific observations.
            with Path(path).open("rb") as f:
                raw = f.read(min(budgets.asset_bytes, 16 * 2**20) + 1)
            if len(raw) > 16 * 2**20:
                raise DawError("over_budget", "metadata parser byte limit")
            parsed = json.loads(raw) if suffix == "json" else ET.fromstring(raw)
            output = {"kind": suffix, "complete": True,
                      "keys": list(parsed)[:100] if isinstance(parsed, dict) else [], "bytes": len(raw)}
        else:
            return {"status": "unsupported_format", "name": name, "reason": f"no enabled reader for .{suffix}"}
        return {"status": "structure_inspected" if output.get("complete") else "partially_inspected", "name": name, **output}
    except ImportError as e:
        return {"status": "unsupported_format", "name": name, "reason": f"optional reader unavailable: {e.name}"}
    except DawError as e:
        return {"status": "unsafe_to_inspect" if e.reason == "unsafe_to_inspect" else "partially_inspected",
                "name": name, "reason": str(e), "complete": False}
    except (ValueError, OSError, KeyError, csv.Error, zipfile.BadZipFile, tarfile.TarError) as e:
        return {"status": "malformed", "name": name, "reason": type(e).__name__ + ": " + str(e)[:300]}


def inspect_asset(ws, aid, isolated=True):
    asset = ws.asset(aid)
    if not asset["blob"]:
        result = {"status": "not_inspected", "reason": asset["access"]}
    else:
        path = ws.blob_path(asset["blob"])
        if file_hash(path) != asset["blob"]:
            raise DawError("integrity_failed", "blob changed since registration")
        if isolated:
            job = ws.root / "staging" / f"inspect-{uuid.uuid4().hex}.json"
            output = job.with_suffix(".out.json")
            write_json(job, {"path": str(path), "name": asset["body"]["name"], "budgets": ws.budgets.model_dump()})
            try:
                proc = subprocess.run([sys.executable, "-m", "daw.worker", str(job), str(output)],
                                      capture_output=True, timeout=ws.budgets.worker_seconds,
                                      env={"PATH": os.environ.get("PATH", ""), "PYTHONPATH": str(Path(__file__).resolve().parents[1]),
                                           "OPENBLAS_NUM_THREADS": "1", "OMP_NUM_THREADS": "1"})
                result = read_json(output) if output.exists() else {
                    "status": "partially_inspected", "reason": "worker_failed", "exit_code": proc.returncode}
            except subprocess.TimeoutExpired:
                result = {"status": "partially_inspected", "reason": "worker_time_budget"}
            finally:
                job.unlink(missing_ok=True)
                output.unlink(missing_ok=True)
        else:
            result = inspect_file(path, asset["body"]["name"], ws.budgets)
    blob = ws.put_json(result)
    iid = "inspect_" + uuid.uuid4().hex
    with ws.db:
        ws.db.execute("INSERT INTO inspection VALUES(?,?,?,?,?)", (iid, aid, result["status"], blob, now()))
    if result.get("kind") in {"archive", "gzip"}:
        chain = asset["body"]["selector"].get("chain", [])
        children = []
        for member in result.get("members", []):
            if member.get("directory"):
                continue
            child = Asset(native_id=f"{aid}!{member['name']}", name=member["name"], size=member.get("size"),
                          license=asset["body"]["license"], selector={"chain": chain + [
                              {"parent": aid, "blob": asset["blob"], "member": member["name"]}]})
            children.append(ws.register_asset(ws.bundle_for(asset), "archive", child, asset["snapshot_id"], previous=aid))
        result["listed_child_assets"] = children
    return {"asset_revision": aid, "inspection": iid, "manifest_blob": blob, **result}


def extract_members(ws, aid, selected):
    asset = ws.asset(aid)
    chain = asset["body"]["selector"].get("chain", [])
    if len(chain) >= ws.budgets.archive_depth:
        raise DawError("over_budget", "archive nesting")
    if not asset["blob"]:
        raise DawError("not_acquired")
    path = ws.blob_path(asset["blob"])
    is_gzip = asset["body"]["name"].endswith(".gz") and not asset["body"]["name"].endswith(".tar.gz")
    if is_gzip:
        inventory = [{"name": asset["body"]["name"][:-3], "directory": False}]
    else:
        inventory = archive_inventory(path, ws.budgets)["members"]
    valid = {m["name"] for m in inventory if not m["directory"]}
    if not selected or not set(selected) <= valid:
        raise DawError("unknown_archive_member")
    results, expanded = [], 0
    for name in dict.fromkeys(selected):
        temp = ws.root / "staging" / (uuid.uuid4().hex + ".member")
        container = None
        try:
            if is_gzip:
                stream = gzip.open(path, "rb")
            elif zipfile.is_zipfile(path):
                container = zipfile.ZipFile(path)
                stream = container.open(name)
            else:
                container = tarfile.open(path)
                stream = container.extractfile(name)
            size = 0
            with stream as inp, temp.open("wb") as out:
                for data in iter(lambda: inp.read(2**20), b""):
                    size += len(data)
                    expanded += len(data)
                    if expanded > ws.budgets.expanded_bytes or size > ws.budgets.asset_bytes:
                        raise DawError("over_budget", "decompressed bytes")
                    ws.check_disk(len(data))
                    out.write(data)
            sha = ws.put_file(temp)
            child = Asset(native_id=f"{aid}!{name}", name=name, size=size, license=asset["body"]["license"],
                          access="available_full", selector={"chain": chain + [{"parent": aid, "blob": asset["blob"], "member": name}]})
            child_id = ws.register_asset(ws.bundle_for(asset), "archive", child, asset["snapshot_id"], sha, previous=aid)
            results.append({"asset_revision": child_id, "member": name, "blob": sha})
        finally:
            if container:
                container.close()
            temp.unlink(missing_ok=True)
    return results
