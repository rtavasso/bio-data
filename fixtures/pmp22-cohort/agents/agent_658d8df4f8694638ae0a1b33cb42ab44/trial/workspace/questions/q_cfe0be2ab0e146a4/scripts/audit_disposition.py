"""Validate primary-source endpoint curation and audit selected co-IP scalars.

No downloaded code, formulas, macros, raw MS or serialized scientific objects run.
This is an endpoint/feature-universe audit, not an independent biological experiment.
"""
import csv
import hashlib
import io
import json
import math
from collections import Counter
from pathlib import Path

import openpyxl
from defusedxml import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs'
OUT.mkdir(exist_ok=True)
manifest = json.loads((ROOT / 'inputs/immutable-manifest.json').read_text())
used = {}
checks = []


def load(name):
    item = manifest[name]
    raw = Path(item['path']).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == item['blob'], name
    used[name] = item['blob']
    return raw


def js(name):
    return json.loads(load(name))


def write_json(name, data):
    raw = json.dumps(data, indent=2, ensure_ascii=False, allow_nan=False) + '\n'
    json.loads(raw, parse_constant=lambda x: (_ for _ in ()).throw(ValueError(x)))
    (OUT / name).write_text(raw)


def write_tsv(name, rows, fields):
    with (OUT / name).open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields, delimiter='\t', extrasaction='raise')
        w.writeheader()
        w.writerows(rows)


curation = js('inputs/curation.json')
design = js('inputs/test-design.json')
roots = {pmc: ET.fromstring(load(f'inputs/public/{pmc}.xml')) for pmc in ['PMC8191293', 'PMC4227013']}
texts = {k: ''.join(v.itertext()) for k, v in roots.items()}
rows = curation['rows']
assert len({r['id'] for r in rows}) == len(rows)
for r in rows:
    assert r['context'] in curation['contexts']
    assert r['anchor'] in texts[r['source']], (r['id'], r['anchor'])
    assert r['function'] == 'not measured'
    assert r['supports_myelin_rescue'] is False
checks.append({'check': 'unique_row_ids_source_anchor_and_context_integrity', 'passed': True, 'rows': len(rows)})

flat = []
for row in rows:
    data = dict(row)
    data.pop('anchor')
    data['source_blob'] = used[f'inputs/public/{row["source"]}.xml']
    data.update(curation['contexts'][row['context']])
    flat.append(data)
write_tsv('disposition-evidence.tsv', flat, list(flat[0]))

wb_raw = load('inputs/public/PMC8191293-mmc2.xlsx')
wb = openpyxl.load_workbook(io.BytesIO(wb_raw), read_only=True, data_only=False, keep_links=False)
assert wb.sheetnames == ['WT', 'N41Q', 'L16P']
expected_header = ('Accession', 'Description', 'Gene', 'Q-value', 'Avg Log2 transformation')
records = []
counts = {}
accessions = {}
nontext = []
formula_cells = []
for ws in wb:
    cells = list(ws.iter_rows())
    assert tuple(c.value for c in cells[0]) == expected_header
    sheetrows = []
    for rownum, cellsrow in enumerate(cells[1:], 2):
        assert len(cellsrow) == 5
        for cell in cellsrow:
            if cell.data_type == 'f':
                formula_cells.append(f'{ws.title}!{cell.coordinate}')
        assert not any(c.data_type in ('f', 'e') for c in cellsrow)
        a, d, g, q, effect = [c.value for c in cellsrow]
        assert isinstance(a, str) and isinstance(d, str)
        assert isinstance(q, (int, float)) and math.isfinite(q) and 0 <= q < 0.1
        assert isinstance(effect, (int, float)) and math.isfinite(effect) and effect > 0.2
        rec = {'sheet': ws.title, 'source_row': rownum, 'accession': a, 'description': d,
               'gene_source_value': g, 'gene_cell_type': cellsrow[2].data_type,
               'gene_eligible_as_symbol': isinstance(g, str), 'q_value': q,
               'mean_log2_ip_over_mock': effect,
               'measurement_status': 'selected_enriched_interactor; not total protein or surface abundance'}
        records.append(rec)
        sheetrows.append(rec)
        if not isinstance(g, str):
            nontext.append({'sheet': ws.title, 'cell': f'C{rownum}', 'accession': a,
                            'literal_value': g, 'cell_type': cellsrow[2].data_type,
                            'treatment': 'preserved, not repaired or used as gene symbol; accession retained'})
    counts[ws.title] = len(sheetrows)
    accessions[ws.title] = {r['accession'] for r in sheetrows}
    assert len(accessions[ws.title]) == len(sheetrows)
wb.close()
assert counts == {'WT': 56, 'N41Q': 485, 'L16P': 129}
assert len(records) == 670
write_tsv('coip-selected-native.tsv', records, list(records[0]))

# Re-read the export and verify every source numeric value was preserved exactly.
with (OUT / 'coip-selected-native.tsv').open() as f:
    exported = list(csv.DictReader(f, delimiter='\t'))
assert len(exported) == len(records)
for original, exported_row in zip(records, exported, strict=True):
    assert original['accession'] == exported_row['accession']
    assert original['source_row'] == int(exported_row['source_row'])
    assert original['q_value'] == float(exported_row['q_value'])
    assert original['mean_log2_ip_over_mock'] == float(exported_row['mean_log2_ip_over_mock'])
checks.append({'check': 'coip_native_numeric_and_accession_roundtrip', 'passed': True,
               'rows': len(records), 'numeric_values': len(records) * 2})

candidates = []
for gene in ['CANX', 'LMAN1', 'RER1', 'UGGT1', 'TMED9', 'ERGIC3', 'SYVN1', 'AMFR']:
    for sheet in wb.sheetnames:
        hits = [r for r in records if r['sheet'] == sheet and r['gene_source_value'] == gene]
        candidates.append({'gene': gene, 'sheet': sheet,
                           'status': 'present_in_selected_export' if hits else 'not_in_selected_export; not measured absence',
                           'native_rows': hits,
                           'functional_interpretation': 'Requires perturbation/localization assay; IP enrichment is neither affinity nor function'})
coip_audit = {
    'question': 'q_cfe0be2ab0e146a4',
    'type': 'Native scalar eligibility audit of an already exposed selected export; not a new regulator screen',
    'source_blob': hashlib.sha256(wb_raw).hexdigest(),
    'counts_by_sheet': counts,
    'selected_rows_total': len(records),
    'unique_accessions_union': len(set.union(*accessions.values())),
    'accessions_common_to_all_sheets': len(set.intersection(*accessions.values())),
    'formula_cells': formula_cells,
    'nontext_gene_cells': nontext,
    'source_filter': 'Q<0.1 and mean log2 IP/mock>0.2; six biological IP replicates per construct across four TMT6 batches; heatmap method additionally requires identification in >=3 replicates',
    'universe': 'Filtered WT/N41Q/L16P lists, not the full measured proteome, six replicate intensities, or perturbation-response universe',
    'candidate_presence': candidates,
    'not_identified': ['Unfiltered measured feature universe', 'Replicate paired total/surface data',
                       'RER1/UGGT1-dependent flux', 'Myelin/conduction function',
                       'Absolute synthesis or degradation rates', 'Variant-difference uncertainty'],
    'interpretation': 'The 670 rows are sheet entries, not 670 distinct regulators. Gene token 5 stays numeric. Missing selected rows are neither measured zeros nor evidence of biological absence. No enrichment recalculation or raw processing performed.'
}
write_json('coip-audit.json', coip_audit)

summaries = []
for r in curation['numeric_summaries']:
    assert 0 < r['control_percent'] <= 100 and 0 <= r['perturbed_percent'] <= 100
    summaries.append({**r, 'percentage_point_difference': r['perturbed_percent'] - r['control_percent'],
                      'ratio_of_reported_summaries': r['perturbed_percent'] / r['control_percent'],
                      'interpretation': 'Descriptive arithmetic only, not mean paired ratio or an inferential estimate; no new CI or p-value'})

inventories = []
for project, filename in [('PXD023091', 'PXD023091-files.json'), ('PXD043917', 'PXD043917-files-current.json')]:
    data = js('inputs/public/' + filename)
    assert isinstance(data, list) and len(data) < 100
    assert all(project in r['projectAccessions'] for r in data)
    inventories.append({'accession': project, 'listed_files': len(data),
                        'requested_page_size': 100, 'category_counts': dict(Counter(r['fileCategory']['value'] for r in data)),
                        'nonraw_files': [{'name': r['fileName'], 'bytes': r['fileSizeBytes'],
                                         'category': r['fileCategory']['value']} for r in data if r['fileCategory']['value'] != 'RAW'],
                        'decision': 'No listed native paired total/surface/myelin export; no raw or search-file processing authorized/performed'})
project = js('inputs/public/PXD043917-project-current.json')
assert '15 days old' in project['sampleProcessingProtocol']
p023091 = js('inputs/public/PXD023091-project.json')
assert 'UGGT1 promotes trafficking' in p023091['projectDescription']
peer = load('peer/intervention-by-endpoint.tsv').decode()
assert 'MG132' in peer and 'MPZ' in peer
load('inputs/prior/REPORT.md')
load('inputs/prior/mechanisms.r021.json')

eligibility = {
    'question': 'q_cfe0be2ab0e146a4', 'agent_authorship': curation['authorship'],
    'contexts': curation['contexts'], 'audited_intervention_contexts': len(rows),
    'qualifying_functional_myelin_rescues_in_audited_rows': sum(r['supports_myelin_rescue'] for r in rows),
    'case_status': [{k: r[k] for k in ['id', 'context', 'locator', 'supports_surface_direction', 'supports_myelin_rescue', 'interpretation']} for r in rows],
    'surface_summary_diagnostics': summaries, 'quantitative_export_inventory': inventories,
    'paired_rer1_total_surface_synthesis_function_matrix': None,
    'missing_status': 'not available in inspected exports; not measured zero; no global claim of nonexistence',
    'pxd043917_blocker': 'Existing source audit: mzIdentML identification scores are not genotype-resolved LFQ. Current PRIDE still describes 15-day nerves; inherited main-figure P7 versus methods/PRIDE P15 conflict remains. No raw reprocessing or assignment of identification scores as protein abundance.',
    'peer_reuse': 'Corrected regulator-turnover table: MG132 restores EGR2 abundance without MPZ rescue. Different cargo/endpoint, so used to reject broad inhibitor as clean functional-rescue instrument, not counted as a PMP22 surface result.',
    'corrections_and_contradictions': [
        'PMC8191293 abstract UGGT1-limiting wording conflicts with Results p0170 and Discussion p0210; PRIDE description agrees with promoting direction.',
        'CNX knockdown HeLa C-terminal GFP and knockout HEK293 extracellular-myc responses are not pooled; no proven explanation of cross-study discrepancy.',
        '2014 RER1 release is L16P-specific versus G150D; endolysosomal localization does not establish rescue.',
        'S2 in 2014 reports Hrd1 siRNA #1 off-target reduction of gp78.',
        '2021 supplementary S8 extracted dose text 10 mM differs from main Fig 4 10 micromolar; do not silently normalize.',
        '2021 S10B caption describes WT/UGGT1 although Results cite it for mutant total-level effects; total-direction statements remain attributed to main text, not independently re-quantified.'
    ]
}
write_json('assay-eligibility.json', eligibility)
write_json('rer1-discriminating-test.json', design)

source_locators = []
for pmc, root in roots.items():
    ids = {el.get('pub-id-type'): el.text for el in root.findall('./front/article-meta/article-id')}
    title = ''.join(root.find('./front/article-meta/title-group/article-title').itertext())
    source_locators.append({'pmcid': pmc, 'pmid': ids['pmid'], 'doi': ids['doi'], 'title': title,
                            'source_blob': used[f'inputs/public/{pmc}.xml'],
                            'row_locators': [{k: r[k] for k in ['id', 'locator']} for r in rows if r['source'] == pmc]})
for name in ['inputs/public/PMC8191293-mmc1.pdf', 'inputs/public/PMC4227013-srep06992-s1.pdf',
             'inputs/public/PMC8191293-mmc3.xlsx']:
    load(name)
# Include only real recorded retrieval receipts. The withdrawn browser gap is not a retrieval.
retrievals = []
for name in manifest:
    if name.endswith('.receipt.json'):
        r = js(name)
        retrievals.append(r)
    elif name.endswith('.receipt.json') is False and '-oa.xml' in name:
        load(name)
write_json('source-locators.json', {'primary_sources': source_locators, 'used_input_blobs': used,
                                  'retrievals': retrievals,
                                  'provenance_correction': 'event_d1ae09b435854271a9575e96a911ed6e browser failure claim was withdrawn by event_224b8f3919cf4172b3ba96bb092e1b6f; no browser evidence used.'})

output_names = ['disposition-evidence.tsv', 'coip-selected-native.tsv', 'coip-audit.json',
                'assay-eligibility.json', 'rer1-discriminating-test.json', 'source-locators.json']
for name in output_names:
    if name.endswith('.json'):
        json.loads((OUT / name).read_text(), parse_constant=lambda x: (_ for _ in ()).throw(ValueError(x)))
checks.append({'check': 'strict_finite_json_roundtrip', 'passed': True})
checks.append({'check': 'native_workbook_no_formulas_errors_or_duplicate_accessions_within_sheet', 'passed': True})
checks.append({'check': 'all_selected_rows_meet_declared_Q_and_log2_thresholds', 'passed': True})
write_json('validation.json', {'valid': True, 'checks': checks,
                               'used_input_blobs': used,
                               'output_hashes': {name: hashlib.sha256((OUT / name).read_bytes()).hexdigest() for name in output_names}})
print(json.dumps({'valid': True, 'contexts': len(rows), 'selected_rows': len(records),
                  'union_accessions': coip_audit['unique_accessions_union'],
                  'shared_accessions': coip_audit['accessions_common_to_all_sheets'],
                  'nontext_gene_cells': nontext, 'output_files': output_names + ['validation.json']}, indent=2))
