"""Inspect complete released metadata and native static supplements before selecting assays."""
from collections import Counter
import csv
import io
import json
from pathlib import Path
from xml.etree import ElementTree as ET
import zipfile
import openpyxl
q=Path(__file__).resolve().parents[1]
o=q/'inputs/views'
files={
'eclip':'doc-cd785330-9240-4913-a1ac-bd765b41b223-attachment-r001',
'hepg2':'doc-3bc11092-2e01-4cd0-b701-ac01a0967f01-attachment-r001',
'k562':'doc-5da1efc3-2a1c-481a-bd17-70b3acb1a65f-attachment-r001',
'secondary':'doc-29dc80a7-b47c-41cb-b759-8dbf5554ebc9-attachment-r001',
'batch':'doc-616f60ab-e379-46a1-823b-fcea881b06bf-attachment-r001'}
allrows=[]
for name,label in files.items():
    with (q/'inputs/http'/(label+'.payload')).open() as f:
        rows=list(csv.DictReader(f,delimiter='\t'))
    print('META',name,len(rows),Counter((r['File format'],r['File output type'],r['Assembly']) for r in rows))
    for row in rows:
        row['publication_collection']=name
        row['metadata_label']=label
        allrows.append(row)
    if name in ['secondary','batch']:
        de=[r for r in rows if 'expression' in r['File output type']]
        print('DE EXAMPLES',json.dumps(de[:2],indent=2))
(q/'inputs/publication-file-metadata.json').write_text(json.dumps(allrows,indent=2,allow_nan=False))
for p in ET.parse(q/'inputs/http/encore-correction-xml-r001.payload').getroot().findall('.//body//p'):
    print('CORRECTION',''.join(p.itertext()))
with zipfile.ZipFile(q/'inputs/http/encore-supplement-zip-r001.payload') as z:
    print('ARCHIVE MEMBERS',[(i.filename,i.file_size) for i in z.infolist()])
    sheets={}
    for member in z.namelist():
        if not member.endswith('.xlsx'):
            continue
        b=z.read(member)
        wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
        sheets[member]={}
        for ws in wb:
            vals=list(ws.values)
            sheets[member][ws.title]=vals
            print('SUPPLEMENT',member,ws.title,'shape',ws.max_row,ws.max_column,'FIRST',repr(vals[:4])[:5500])
        wb.close()
    (o/'supplement-tables.json').write_text(json.dumps(sheets,indent=2,allow_nan=False,default=str))
