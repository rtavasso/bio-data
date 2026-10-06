"""Safely inspect acquired XML and selected numeric source containers."""
import hashlib
import io
import json
import zipfile
from pathlib import Path
from defusedxml import ElementTree as ET
import openpyxl
q = Path(__file__).resolve().parents[1]
ws = q.parents[1]
manifest = []
for receipt in sorted((q/'outputs').glob('*-fetch.json')):
    r = json.loads(receipt.read_text())
    if not r.get('blob'):
        continue
    p = ws/'blobs/sha256'/r['blob'][:2]/r['blob']
    data = p.read_bytes()
    assert hashlib.sha256(data).hexdigest() == r['blob']
    if receipt.name.endswith('.xml-fetch.json'):
        root = ET.fromstring(data)
        lines = []
        for i, e in enumerate(root.iter()):
            if e.tag in ['article-title','p','fig','table-wrap','ref','supplementary-material']:
                text = ' '.join(' '.join(e.itertext()).split())
                lines.append(f'element={i} tag={e.tag} id={e.get("id")} {text}')
        out = q/'outputs'/receipt.name.replace('.xml-fetch.json','.source.txt')
        out.write_text('\n'.join(lines)+'\n')
        print('XML', receipt.name, r['blob'], 'paragraphs', len(lines))
        for line in lines:
            if 'GSE' in line or 'data availability' in line.lower():
                print(line)
        manifest.append({'receipt':str(receipt.relative_to(q)), 'blob':r['blob'],'output':str(out.relative_to(q))})
    elif receipt.name.endswith('.zip-fetch.json'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            print('ZIP', receipt.name, [(i.filename,i.file_size) for i in z.infolist()])
            for info in z.infolist():
                if info.filename.endswith(('.xlsx','.csv','.txt','.pzfx')) and not info.filename.startswith('__MACOSX'):
                    b = z.read(info.filename)
                    out = q/'inputs'/'tead-numeric'/Path(info.filename).name
                    out.parent.mkdir(parents=True,exist_ok=True)
                    assert not out.exists()
                    out.write_bytes(b)
                    item = {'receipt':str(receipt.relative_to(q)),'container_blob':r['blob'],'member':info.filename,'sha256':hashlib.sha256(b).hexdigest(),'output':str(out.relative_to(q))}
                    manifest.append(item)
                    print('EXTRACTED', item)
                    if out.suffix == '.xlsx':
                        wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=False)
                        for sheet in wb:
                            print('SHEET',sheet.title,sheet.max_row,sheet.max_column)
                            for row in sheet.iter_rows(min_row=1,max_row=min(12,sheet.max_row),values_only=True):
                                print(repr(row))
(q/'outputs/acquired-source-inspection.json').write_text(json.dumps(manifest,indent=2)+'\n')
