"""Inspect new source units and endpoint locators, without selecting on target effect."""
import json
import re
from pathlib import Path
import openpyxl
q=Path(__file__).resolve().parents[1]
ws=q.parents[1]
r=json.loads((q/'outputs/nedd4-tableS1-fetch.json').read_text())
p=ws/'blobs/sha256'/r['blob'][:2]/r['blob']
with p.open('rb') as f:
    wb=openpyxl.load_workbook(f,read_only=True,data_only=False)
    for s in wb:
        print('NEDD4_SHEET',s.title,s.max_row,s.max_column)
        for row in s.iter_rows(min_row=1,max_row=6,values_only=True):
            print(repr(row))
for pmc in ['PMC7498331','PMC5589416','PMC3100536']:
    d=json.loads((q/f'outputs/{pmc}.text.json').read_text())
    print('ACCESSIONS',pmc,sorted(set(re.findall(r'GSE\d+',str(d)))))
    for r in d:
        text=r['text']
        if pmc=='PMC7498331' and any(t in text for t in ['Total RNA','transcriptome data','deposited','Egr2 expression','expression of Egr2','Egr2 levels','Sox10 and Egr2']):
            print('RNF40',r)
        if pmc=='PMC5589416' and 'Given our focus' in text:
            print('NOVELTY',r)
        if pmc=='PMC3100536' and any(t in text for t in ['Prx','Mag','Mpz','Mbp']) and r['tag']=='p':
            print('PROGRAM_SOURCE',r)
