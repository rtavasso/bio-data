"""Validate output bytes against the producing receipt and original workbook."""
import csv
import hashlib
import io
import json
from pathlib import Path

import openpyxl

P = Path(__file__).resolve().parents[1]
OUT = P / 'outputs'
receipt = json.loads((OUT / 'execution-r002.json').read_text())
assert receipt['complete'] and receipt['exit_code'] == 0 and receipt['code_unchanged']
assert hashlib.sha256(Path(receipt['producer']).read_bytes()).hexdigest() == receipt['code_sha256']
for item in receipt['outputs']:
    assert item['written']
    assert hashlib.sha256(Path(item['path']).read_bytes()).hexdigest() == item['sha256']


def reject_constant(value):
    raise ValueError(value)


bundle = json.loads((OUT / 'evidence-package.json').read_text(), parse_constant=reject_constant)
assert bundle['checks'] and all(bundle['checks'].values())
manifest = json.loads((P / 'sources/immutable-inputs.json').read_text())
x = Path(manifest['primary/12864_2020_6963_MOESM2_ESM.xlsx']['path']).read_bytes()
w = openpyxl.load_workbook(io.BytesIO(x), read_only=True, data_only=False)
s = w['TSS- H3K4me3 and SOX10']
expected = {}
for i, row in enumerate(s.iter_rows(values_only=True), 1):
    if i > 2 and row[4] == 'Pmp22':
        expected[i] = row
actual = bundle['tss_audit']['pmp22_clusters']
header = list(next(s.iter_rows(min_row=2, max_row=2, values_only=True)))
assert set(expected) == {r['worksheet_row'] for r in actual}
for r in actual:
    assert tuple(r[k] for k in header) == expected[r['worksheet_row']]
assert len(actual) == 22
for name, n in [('cis-evidence.tsv', 12), ('rat-promoter-map.tsv', 2),
                ('pmp22-tss-clusters.tsv', 22), ('sample-eligibility.tsv', 14)]:
    with (OUT / name).open() as f:
        rows = list(csv.DictReader(f, delimiter='\t'))
    assert len(rows) == n, (name, len(rows))
assert sum(r['eligible_for_nascent_cis_test'] for r in bundle['assay_eligibility']['samples']) == 0
assert bundle['tss_audit']['per_library_data_available'] is False
assert bundle['promoter_mapping']['human_endogenous_P1_P2_coordinates'] is None
assert len({r['id'] for r in bundle['agent_curated_evidence']['rows']}) == 12
report = {'status': 'passed', 'execution_receipt_complete': True, 'output_hashes_match': True,
          'native_workbook_rows_exact': 22, 'matrix_rows': 12, 'rat_primer_anchors': 2,
          'sample_rows': 14, 'strict_JSON_finite': True,
          'scope': 'Data-integrity validation, not independent scientific replication',
          'execution_receipt_sha256': hashlib.sha256((OUT / 'execution-r002.json').read_bytes()).hexdigest()}
(OUT / 'validation.json').write_text(json.dumps(report, indent=2, allow_nan=False) + '\n')
print(json.dumps(report))
