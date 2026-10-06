"""Bounded, retrospective extraction and exact-locus mapping of immutable sources.

This producer calculates no new differential-expression statistics. Evidence
interpretations are explicitly attributed to the separate agent-curated input.
"""
import csv
import hashlib
import io
import json
import math
import platform
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs'
OUT.mkdir(exist_ok=True)
inputs = json.loads((ROOT / 'sources/immutable-inputs.json').read_text())


def raw(name):
    info = inputs[name]
    p = Path(info['path'])
    b = p.read_bytes()
    assert hashlib.sha256(b).hexdigest() == info['blob'], name
    return b


def obj(name):
    return json.loads(raw(name))


def tsv(name, rows):
    assert rows
    keys = list(rows[0])
    with (OUT / name).open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=keys, delimiter='\t')
        w.writeheader()
        for row in rows:
            assert list(row) == keys
            w.writerow({k: ('NA' if v is None else json.dumps(v, ensure_ascii=False)
                            if isinstance(v, (dict, list)) else v) for k, v in row.items()})


# Validate every declared immutable input, including native XML/PDF and receipts.
for name in inputs:
    raw(name)
curation = obj('curated-evidence.json')
assert len({r['id'] for r in curation['rows']}) == len(curation['rows'])
for r in curation['rows']:
    assert r['source_file'] in inputs
    r['source_blob'] = inputs[r['source_file']]['blob']

book = openpyxl.load_workbook(io.BytesIO(raw('primary/12864_2020_6963_MOESM2_ESM.xlsx')),
                              read_only=True, data_only=False)
sheet_audit = []
all_rows = []
formula_count = 0
for s in book:
    header = [c.value for c in next(s.iter_rows(min_row=2, max_row=2))]
    valid = []
    for number, cells in enumerate(s.iter_rows(min_row=3), 3):
        formula_count += sum(c.data_type == 'f' for c in cells)
        row = [c.value for c in cells]
        if all(v is None for v in row):
            continue
        assert isinstance(row[3], int), (s.title, number, row[:6])
        valid.append(dict(zip(header, row, strict=True)))
        if s == book.worksheets[0]:
            all_rows.append({'source_sheet': s.title, 'worksheet_row': number, **valid[-1]})
    sheet_audit.append({'sheet': s.title, 'declared_max_row': s.max_row,
                        'nonempty_data_rows': len(valid), 'unique_cluster_ids': len({r['TSS id'] for r in valid})})
assert formula_count == 0, 'Never evaluate downloaded formulas'
assert len(all_rows) == 4993
assert len({r['TSS id'] for r in all_rows}) == 4993
pmp = [r for r in all_rows if r['Gene Name'] == 'Pmp22']
assert len(pmp) == 22
assert {r['TSS Chr'] for r in pmp} == {'chr10'}
assert {r['TSS Strand'] for r in pmp} == {'+'}
by_id = {r['TSS id']: r for r in pmp}
for cid, expected in {5439: (49316968, 49317054), 5446: (49319494, 49319573)}.items():
    assert (by_id[cid]['TSS Start'], by_id[cid]['TSS End']) == expected

# Metadata discovery is recursive because bio's source receipts nest native GEO fields.
samples = {}


def walk(x):
    if isinstance(x, dict):
        f = x.get('fields', {})
        if f.get('Sample_geo_accession'):
            acc = f['Sample_geo_accession'][0]
            rec = {k: f.get(k) for k in [
                'Sample_geo_accession', 'Sample_title', 'Sample_source_name_ch1',
                'Sample_characteristics_ch1', 'Sample_library_strategy',
                'Sample_library_source', 'Sample_library_selection', 'Sample_data_processing']}
            if acc in samples:
                assert samples[acc] == rec
            samples[acc] = rec
        for v in x.values():
            walk(v)
    elif isinstance(x, list):
        for v in x:
            walk(v)


walk(obj('GSE139321-metadata.json'))
assert len(samples) == 14
sample_map = []
for acc, r in sorted(samples.items()):
    title = r['Sample_title'][0]
    if 'Sciatic Nerve' in title:
        group = 'adult_nerve'
        unit = 'Independent nerve RNA sample; animal pairing/identity not established'
        time = '6-9 months old; no intervention'
    elif 'Primary' in title:
        group = 'primary_cAMP' if 'cAMP' in title else 'primary_control'
        unit = 'Independent culture population; commercial primary cells; independent donors not established'
        time = 'CPT-cAMP or vehicle from culture day4 through day7'
    else:
        group = 'S16_parental' if 'Parental' in title else 'S16_SOX10_KO'
        unit = 'Parental independent RNA sample' if 'Parental' in title else 'Single-cell-derived clone; four clones across two guides, not four donors'
        time = 'Established parental/knockout line; no acute interval established'
    sample_map.append({'accession': acc, 'title': title, 'group': group, 'unit': unit,
                       'time': time, 'assembly': 'rn5', 'assay': 'Tn5Prime on total RNA',
                       'eligible_for_nascent_cis_test': False, 'native_fields': r})
counts = {g: sum(r['group'] == g for r in sample_map) for g in sorted({r['group'] for r in sample_map})}
assert counts == {'S16_SOX10_KO': 4, 'S16_parental': 2, 'adult_nerve': 2, 'primary_cAMP': 3, 'primary_control': 3}

# Independently anchor published first-exon primers to rat genomic DNA.
seq = obj('primary/rn5-Pmp22-promoter-sequence.json')
dna = seq['dna'].upper()
assert seq['genome'] == 'rn5' and seq['chrom'] == 'chr10'
assert len(dna) == seq['end'] - seq['start']
thesis = raw('primary/Pantera-thesis.txt').decode()
table2 = thesis.split('[PDF page 78]', 1)[1].split('[PDF page 79]', 1)[0]
assert 'Table 2-1. Primers used for qRT-PCR.' in table2
primers = re.findall(r'^Pmp22-(P[12]) ([ACGT]+) ([ACGT]+)$', table2, re.M)
assert len(primers) == 2
maps = []
for promoter, forward, reverse in primers:
    hits = [m.start() for m in re.finditer('(?=' + forward + ')', dna)]
    assert len(hits) == 1, (promoter, hits)
    start = seq['start'] + hits[0]
    cluster = by_id[5439 if promoter == 'P1' else 5446]
    maps.append({'species': 'Rattus norvegicus', 'assembly': 'rn5', 'chrom': 'chr10', 'strand': '+',
                 'promoter_label': promoter + '-associated', 'first_exon_label': '1A' if promoter == 'P1' else '1B',
                 'cluster_id': cluster['TSS id'], 'cluster_start_source': cluster['TSS Start'],
                 'cluster_end_source': cluster['TSS End'],
                 'cluster_coordinate_convention': 'Source tokens preserved; no BED conversion asserted',
                 'forward_primer': forward, 'reverse_primer_source_not_mapped': reverse,
                 'primer_start_0based': start, 'primer_end_exclusive_0based': start + len(forward),
                 'primer_exact_hits_in_queried_interval': len(hits),
                 'mapping_support': 'Exact source promoter-specific forward-primer anchor near start cluster; no genome-wide uniqueness/cap/full-transcript linkage claim',
                 'primer_locator': 'Pantera thesis printed p70 / PDF78 Table2-1',
                 'human_lifted_coordinates': None})
ref = obj('primary/rn5-refGene-Pmp22.json')
assert len(ref['refGene']) == 1
transcript = ref['refGene'][0]
assert transcript['name'] == 'NM_017037' and transcript['name2'] == 'Pmp22'
assert by_id[5439]['TSS Start'] <= transcript['txStart'] < by_id[5439]['TSS End']

# Preserve exact historical hg18 element coordinates, never liftover by position.
bioc = ET.fromstring(raw('primary/PMC3100536-bioc.xml'))
passages = [''.join(e.itertext()) for e in bioc.iter('text')]
assert any('15,091,959-15,092,201' in p for p in passages)
assert any('15,090,965-15,092,611' in p for p in passages)
other_map = [
    {'species': 'Homo sapiens', 'assembly_source': 'hg18 Mar.2006', 'chrom': 'chr17',
     'element': '+11kb intronic reporter', 'start_source': 15091959, 'end_source': 15092201,
     'coordinate_convention': 'Source printed interval; no conversion',
     'endpoint': 'Human sequence tested in mouse B16/F10 cells, not human endogenous assay',
     'locator': 'PMC3100536 BioC passage15'},
    {'species': 'Homo sapiens', 'assembly_source': 'hg18 Mar.2006', 'chrom': 'chr17',
     'element': 'Intronic transgene fragment', 'start_source': 15090965, 'end_source': 15092611,
     'coordinate_convention': 'Source printed interval; no conversion',
     'endpoint': 'Hsp68 transgene in mice, not endogenous P1/P2 promoter',
     'locator': 'PMC3100536 BioC passage22'}]
for r in pmp:
    for k, v in r.items():
        if isinstance(v, float):
            assert math.isfinite(v), (k, v)

state_rows = [r for r in all_rows if r['Gene Name'] in {'Mpz', 'Mag', 'Egr2', 'Sox10'}]
transport = obj('transport-receipts.json')
for r in transport:
    assert hashlib.sha256(r['original_text'].encode()).hexdigest() == r['sha256']

result = {
    'question': curation['question'], 'status': 'bounded audit complete; causal initiation contrast not identifiable',
    'analysis_type': 'Retrospective source verification and local primer mapping; no independent biological replication or new differential test',
    'environment': {'python': sys.version, 'platform': platform.platform(), 'openpyxl': openpyxl.__version__},
    'input_objects': {n: {'blob': i['blob'], 'bytes': i['bytes']} for n, i in inputs.items()},
    'agent_curated_evidence': curation,
    'assay_eligibility': {'samples': sample_map, 'group_counts': counts,
                        'causal_test_eligible': False,
                        'reason': 'Trans/state perturbations of total-RNA starts; no matched cis intervention, nascent labelling, decay or early time course'},
    'tss_audit': {'selected_universe_rows': len(all_rows), 'Pmp22_rows': len(pmp),
                  'workbook_sheets': sheet_audit, 'formulas_encountered': formula_count,
                  'feature_universe': '4993 adult-nerve TSSs associated with H3K4me3 and SOX10; not every Pmp22 start',
                  'units': 'Group mean reads per million and author-provided edgeR log2FC/FDR, not raw counts',
                  'blank_semantics': 'Source blank -> JSON null / TSV NA; test not supplied, not a zero effect',
                  'zero_semantics': 'Numeric zero retained as reported group RPM; not universal biological absence',
                  'per_library_data_available': False,
                  'pmp22_clusters': pmp, 'state_marker_rows_source_only': state_rows,
                  'focal_clusters': [by_id[5439], by_id[5446]],
                  'statistical_limits': 'No new CI, promoter-interaction p-value or ratio-of-means differential test; source group means cannot recover replicate covariance'},
    'promoter_mapping': {'rat_local_anchors': maps, 'rn5_refGene_record': transcript,
                        'current_RefSeq_caution': curation['coordinate_cautions'][3],
                        'source_defined_human_reporter_elements': other_map,
                        'human_endogenous_P1_P2_coordinates': None,
                        'unresolved': 'No human promoter coordinate transfer; no validated full first-exon connection for every minor cluster'},
    'transport_receipts': transport,
    'checks': {'input_hashes': True, 'sample_count_and_groups': True, 'workbook_unique_ids': True,
               'no_formula_evaluation': True, 'exact_forward_primer_hits': True,
               'nonfinite_numbers_absent': True}
}
(OUT / 'evidence-package.json').write_text(json.dumps(result, indent=2, ensure_ascii=False, allow_nan=False) + '\n')
tsv('cis-evidence.tsv', curation['rows'])
tsv('rat-promoter-map.tsv', maps)
tsv('pmp22-tss-clusters.tsv', pmp)
tsv('sample-eligibility.tsv', sample_map)
print(json.dumps({'checks': result['checks'], 'matrix_rows': len(curation['rows']),
                  'tss_universe': len(all_rows), 'Pmp22_rows': len(pmp), 'samples': len(sample_map),
                  'sheet_audit': sheet_audit, 'primer_anchors': maps,
                  'focal_source_results': [{k: r[k] for k in ['TSS id', 'worksheet_row', 'Primary SC cAMP vs Control log2FC', 'S16 ΔSOX10 vs Parental log2FC']} for r in result['tss_audit']['focal_clusters']]}, ensure_ascii=False, indent=2))
