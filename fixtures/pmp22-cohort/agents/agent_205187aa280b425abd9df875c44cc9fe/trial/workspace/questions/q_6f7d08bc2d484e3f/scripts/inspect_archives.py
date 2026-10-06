"""Inventory and selectively extract nonexecutables from source ZIPs."""
from pathlib import Path
import json
import zipfile
import hashlib
import openpyxl
Q = Path(__file__).resolve().parents[1]
manifest=[]
for accession in ('PMC10684506','PMC10316758','PMC11370802'):
    p=Q/'inputs/public'/f'{accession}-supp.zip'
    with zipfile.ZipFile(p) as z:
        for i in z.infolist():
            if i.filename.endswith(('.pdf','.xlsx','.zip','.docx')):
                print(accession,i.filename,i.file_size)
                if i.filename.endswith(('.xlsx','.zip','.docx')) or 'MOESM5' in i.filename:
                    dest=Q/'inputs/native'/accession/Path(i.filename).name
                    dest.parent.mkdir(parents=True,exist_ok=True)
                    data=z.read(i)
                    if dest.exists(): assert dest.read_bytes()==data
                    else: dest.write_bytes(data)
                    manifest.append({'archive':str(p.relative_to(Q)),'archive_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'member':i.filename,'path':str(dest.relative_to(Q)),'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data)})
                    if dest.suffix=='.zip':
                        with zipfile.ZipFile(dest) as nested:
                            for j in nested.infolist():
                                print('NESTED',j.filename,j.file_size)
                                if j.filename.endswith(('.pdf','.docx','.xlsx')) and not j.is_dir():
                                    d=Q/'inputs/native'/accession/Path(j.filename).name
                                    b=nested.read(j)
                                    d.write_bytes(b)
                                    manifest.append({'archive':str(dest.relative_to(Q)),'member':j.filename,'path':str(d.relative_to(Q)),'sha256':hashlib.sha256(b).hexdigest(),'bytes':len(b)})
(Q/'inputs/native/extraction-manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
p=Q/'inputs/native/PMC10684506/43856_2023_400_MOESM2_ESM.xlsx'
w=openpyxl.load_workbook(p,read_only=True,data_only=False)
lines=[]
for s in w:
    lines.append(f'SHEET {s.title!r} rows={s.max_row} cols={s.max_column}')
    for row in s:
        cells=[f'{c.coordinate}={c.value!r} ({c.data_type})' for c in row if c.value is not None]
        if cells: lines.append(' | '.join(cells))
(Q/'inputs/text/yoshioka-workbook.txt').write_text('\n'.join(lines)+'\n')
print('Workbook sheets',w.sheetnames)
w.close()
