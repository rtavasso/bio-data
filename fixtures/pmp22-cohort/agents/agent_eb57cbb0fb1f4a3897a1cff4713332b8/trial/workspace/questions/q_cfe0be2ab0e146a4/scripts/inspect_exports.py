"""Safe scalar inspection of declared native workbooks and provider files."""
import json
from pathlib import Path
import openpyxl
from defusedxml import ElementTree as ET
ROOT=Path(__file__).resolve().parents[1]
for name in ['PXD023091-files.json','PXD043917-files-current.json']:
    doc=json.loads((ROOT/'inputs/public'/name).read_text())
    print(name,'returned',len(doc))
    for f in doc:
        if f['fileCategory']['value']!='RAW':
            print(f['fileName'],f['fileSizeBytes'],f['fileCategory']['value'],[x['value'] for x in f['publicFileLocations'] if x['name']=='FTP Protocol'])
for name in ['PMC8191293-mmc2.xlsx','PMC8191293-mmc3.xlsx']:
    wb=openpyxl.load_workbook(ROOT/'inputs/public'/name,read_only=True,data_only=False,keep_links=False)
    print(name,wb.sheetnames)
    for ws in wb:
        print(ws.title,ws.max_row,ws.max_column)
        for idx,row in enumerate(ws.iter_rows(),1):
            if idx<=3 or any(c.data_type=='f' or (c.value is not None and c.column==3 and not isinstance(c.value,str)) for c in row):
                print(idx,[(c.coordinate,c.value,c.data_type) for c in row])
    wb.close()
for pmc in ['PMC8191293','PMC4227013']:
    root=ET.parse(ROOT/'inputs/public'/f'{pmc}.xml').getroot()
    print(pmc,[(e.attrib,e.text) for e in root.findall('./front/article-meta/article-id')])
