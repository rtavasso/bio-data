"""Reviewed literal/table recipe for the exact acquired Schwann-cell pilot files.

No sample normalization, promoter assignment, cross-species mapping, or statistics.
Every semantic setting cites an original workbook cell/header or saved GEO SOFT.
"""
import argparse
import csv

import openpyxl

from daw.catalog import Workspace
from daw.curation import accept, add_assertion
from daw.models import AssertionInput, Curation, Evidence, Query
from daw.query import query, render_report
from daw.util import DawError, write_json


EXPECTED = {
    "12864_2020_6963_MOESM2_ESM.xlsx": "3534c0aa574c7250774426cc248212fcea9b3f402b6338aa08d3ffc491668c4f",
    "12864_2020_6963_MOESM3_ESM.xlsx": "309e623447b735c327f1478bf22aacd06ca41852d81430952720a4d354c0c01b",
    "GSE139321_Schwann_Cell_Tn5Prime_GEO_Processed.txt": "fe9f2e05e9315040f4237b130b0e53d4bdc68b2ca9b4fc52a22e4f5520e5f7d5",
    "GSE201623_LentiAS.rnaseq.counts.txt": "b3d15fe88a6a680290cb57a382ac93c9f5451966d5d20b58bfa2994b0ddcb580",
}


def assertion(ws, asset, field, value, raw, locator, selector=None, evidence_blob=None, method="reviewed literal source mapping"):
    return add_assertion(ws, AssertionInput(subject=asset["resource_id"], field=field, raw_value=raw, value=value,
        evidence=Evidence(blob=evidence_blob or asset["blob"], locator=locator, method=method), scope=selector or {}))


def geo_evidence(ws, accession, contains):
    rows = ws.rows("SELECT blob FROM snapshot WHERE locator LIKE ? AND outcome='available_full' ORDER BY retrieved DESC",
                   (f"%targ=gsm&acc={accession}&%",))
    for row in rows:
        for line in ws.blob_path(row["blob"]).read_text().splitlines():
            if contains in line:
                return row["blob"], line
    raise DawError("recipe_source_evidence_missing", accession + ": " + contains)


def run(ws):
    assets = {a["body"]["name"]: a for a in ws.assets() if a["blob"] and a["body"]["name"] in EXPECTED}
    for name, asset in assets.items():
        if asset["blob"] != EXPECTED[name]:
            raise DawError("recipe_input_changed", name + ": re-review original bytes before applying this recipe")
    accepted = []
    for name, asset in assets.items():
        if name.endswith("MOESM2_ESM.xlsx"):
            book = openpyxl.load_workbook(ws.blob_path(asset["blob"]).open("rb"), read_only=True, data_only=False)
            try:
                for sheet_index, sheet in enumerate(book):
                    title, headers = sheet["A1"].value, [c.value for c in sheet[2]]
                    assert headers[4] == "Gene Name"
                    for contrast_name, columns, contrast in [
                        ("cAMP_vs_control", {"cAMP_log2FC": 16, "cAMP_FDR": 17},
                         {"numerator": "cAMP-treated primary rat Schwann cells", "denominator": "control primary rat Schwann cells", "effect_type": "published_edgeR_log2_fold_change"}),
                        ("SOX10_loss_vs_parental", {"SOX10_loss_log2FC": 20, "SOX10_loss_FDR": 21},
                         {"numerator": "SOX10-deleted rat S16 cells", "denominator": "parental rat S16 cells", "effect_type": "published_edgeR_log2_fold_change"}),
                    ]:
                        selector = {"sheet": sheet.title, "start_row": 3, "end_row": sheet.max_row, "contrast": contrast_name}
                        selection = "selected" if sheet_index == 0 else "significant_only"
                        units = {key: "log2 fold change" if key.endswith("log2FC") else "FDR-corrected p-value" for key in columns}
                        settings = {"feature_column": 5, "columns": columns, "namespace": "literal_rat_source_gene_name",
                                    "selection": selection, "units": "per_field", "column_units": units, "contrast": contrast}
                        assertions = {}
                        for field in ("namespace", "selection", "units", "column_units", "contrast"):
                            raw = "Gene Name" if field == "namespace" else title if field == "selection" else [headers]
                            locator = f"sheet:{sheet.title}!" + ("E2" if field == "namespace" else "A1" if field == "selection" else "A2:U2")
                            assertions[field] = assertion(ws, asset, field, settings[field], raw, locator, selector)
                        proposal = Curation(asset_revision=asset["id"], kind="table", selector=selector, settings=settings,
                            assertions=assertions, limitations=["Rat TSS-level published estimates; not gene-wide or P1/P2 estimates",
                            "Restricted to source-selected SOX10/H3K4me3-associated TSSs; rows and contrasts are not independent experiments"])
                        accepted.append(accept(ws, proposal))
            finally:
                book.close()
        elif name.endswith("MOESM3_ESM.xlsx"):
            book = openpyxl.load_workbook(ws.blob_path(asset["blob"]).open("rb"), read_only=True, data_only=False)
            try:
                sheet = book["Sheet 1"]
                for col in range(1, 5):
                    selector = {"sheet": "Sheet 1", "start_row": 3, "end_row": sheet.max_row, "list_column": col}
                    settings = {"feature_column": col, "columns": {}, "namespace": "literal_source_gene_label",
                                "selection": "selected"}
                    assertions = {f: assertion(ws, asset, f, value, sheet["A1"].value, "sheet:Sheet 1!A1", selector)
                                  for f, value in settings.items() if f in {"namespace", "selection"}}
                    accepted.append(accept(ws, Curation(asset_revision=asset["id"], kind="table", selector=selector,
                        settings=settings, assertions=assertions,
                        limitations=["Selected source list linking rat TSS experiments to human RefSeq isoform annotation; no cross-species activity equivalence"])))
            finally:
                book.close()
        else:
            with ws.blob_path(asset["blob"]).open() as f:
                first = f.readline().rstrip("\r\n")
            headers = next(csv.reader([first], delimiter="\t"))
            if name.startswith("GSE139321"):
                settings = {"feature_column": 5, "columns": {headers[c - 1]: c for c in (13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 29, 30)},
                            "namespace": "literal_rat_source_gene_name", "selection": "selected", "units": "per_field",
                            "column_units": {headers[c - 1]: "RPM" if "RPM" in headers[c - 1] else "FDR-corrected p-value" if "P-value" in headers[c - 1] else "log2 fold change" for c in (13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 29, 30)}}
            else:
                settings = {"feature_column": 1, "columns": {headers[c - 1]: c for c in (2, 3, 4, 5)},
                            "namespace": "literal_rat_source_gene_name", "selection": "unknown", "units": "source_reported_counts"}
            assertions = {}
            for field, value in settings.items():
                if field not in {"namespace", "selection", "units", "column_units"}:
                    continue
                blob, raw, locator = asset["blob"], first, "lines:1-1"
                if name.startswith("GSE139321") and field == "selection":
                    blob, raw = geo_evidence(ws, "GSE139321", "4,993 TSSs")
                    locator = "text:" + raw
                if name.startswith("GSE201623") and field == "units":
                    blob, raw = geo_evidence(ws, "GSE201623", "Read summarization was done with featureCounts")
                    locator = "text:" + raw
                assertions[field] = assertion(ws, asset, field, value, raw, locator, evidence_blob=blob,
                    method="source-backed conservative mapping; unknown completeness is not an asserted tested universe")
            # Selection is intentionally conservative: the header does not certify a complete tested universe.
            accepted.append(accept(ws, Curation(asset_revision=asset["id"], kind="table", settings=settings, assertions=assertions,
                limitations=["Literal source values; no new differential test or donor inference", "Native study-specific units remain separate"])))
    scope = sorted({ws.bundle_for(a) for a in assets.values()})
    output = []
    for feature in ("Pmp22", "Sox10"):
        result = query(ws, Query(question_id=f"Rat {feature} · source-reported Schwann measurements", operator="table.feature_lookup", feature=feature, scope=scope))
        output.append({**result, **render_report(ws, result["run"])})
    output.append(query(ws, Query(question_id="Pmp22 · published TSS contrasts with restricted selection", operator="contrast.published_lookup", feature="Pmp22", scope=scope)))
    output.append(query(ws, Query(question_id="PMP22 · auxiliary selected isoform lists", operator="table.feature_lookup", feature="PMP22", scope=scope)))
    output.append(query(ws, Query(question_id="Negative control · selected-list omission cannot establish null", operator="contrast.published_lookup", feature="DAW_WITHHELD_ABSENT_IDENTIFIER", scope=scope)))
    write_json("docs/receipts/pilot-table-results.json", {"recipe": "pilot_tables.py", "inputs": EXPECTED,
               "accepted_curations": accepted, "queries": output, "scope": scope})
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", default="workspaces/pilot")
    args = parser.parse_args()
    workspace = Workspace(args.workspace)
    with workspace.writer():
        for result in run(workspace):
            print(result["run"], result["measurements"], result.get("report", ""), flush=True)
    workspace.close()
