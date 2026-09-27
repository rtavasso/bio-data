import io
import tarfile
import zipfile

import openpyxl
import pytest

from daw.inspectors import archive_inventory, extract_members, inspect_asset, inspect_file
from daw.models import Budgets
from daw.util import DawError


def test_hidden_sheet_formula_offsets_and_literal_na(ws, tmp_path):
    path = tmp_path / "supp.xlsx"
    book = openpyxl.Workbook()
    hidden = book.create_sheet("hidden auxiliary")
    hidden.sheet_state = "hidden"
    hidden["C5"] = "NA"
    hidden["D5"] = "=1+1"
    book.active.merge_cells("A1:C1")
    book.save(path)
    result = inspect_asset(ws, ws.local_asset(path)["asset_revision"])
    assert result["status"] == "structure_inspected"
    sheet = result["sheets"][1]
    assert sheet["state"] == "hidden"
    assert sheet["preview"][0][0] == {"cell": "C5", "value": "NA", "type": "s"}
    assert sheet["formulas"][0]["cached"] is None
    assert result["sheets"][0]["merged_ranges"] == ["A1:C1"]


@pytest.mark.parametrize("name", ["../escape", "/tmp/escape", "dir/../../escape", "C:\\escape"])
def test_zip_traversal_rejected(tmp_path, name):
    path = tmp_path / "archive.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(name, "bad")
    with pytest.raises(DawError, match="unsafe"):
        archive_inventory(path, Budgets())


def test_duplicate_members_and_expansion(tmp_path):
    path = tmp_path / "archive.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("x", "x" * 100)
    with pytest.raises(DawError, match="over_budget"):
        archive_inventory(path, Budgets(expanded_bytes=50))
    with zipfile.ZipFile(path, "a") as z, pytest.warns(UserWarning):
        z.writestr("x", "duplicate")
    with pytest.raises(DawError, match="unsafe"):
        archive_inventory(path, Budgets())


def test_tar_symlinks_rejected(tmp_path):
    path = tmp_path / "archive.tar"
    with tarfile.open(path, "w") as tar:
        info = tarfile.TarInfo("link")
        info.type, info.linkname = tarfile.SYMTYPE, "/etc/passwd"
        tar.addfile(info)
    with pytest.raises(DawError, match="unsafe"):
        archive_inventory(path, Budgets())


def test_nested_archive_parent_chain(ws, tmp_path):
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as z:
        z.writestr("table.tsv", "gene\tvalue\nNA\t0\n")
    path = tmp_path / "outer.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("nested.zip", inner.getvalue())
    outer = ws.local_asset(path)
    first = extract_members(ws, outer["asset_revision"], ["nested.zip"])[0]
    second = extract_members(ws, first["asset_revision"], ["table.tsv"])[0]
    chain = ws.asset(second["asset_revision"])["body"]["selector"]["chain"]
    assert len(chain) == 2 and chain[0]["blob"] == outer["blob"] and chain[1]["blob"] == first["blob"]


def test_pickle_source_is_never_executed(tmp_path):
    path = tmp_path / "malicious.pkl"
    path.write_bytes(b"not executed")
    assert inspect_file(path, path.name)["status"] == "unsupported_format"


def test_hdf_external_links_rejected(tmp_path):
    h5py = pytest.importorskip("h5py")
    path = tmp_path / "external.h5ad"
    with h5py.File(path, "w") as f:
        f["outside"] = h5py.ExternalLink("/etc/passwd", "/")
    assert inspect_file(path, path.name)["status"] == "unsafe_to_inspect"
